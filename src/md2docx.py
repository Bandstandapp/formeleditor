"""ANLEITUNG.md → docs/Bedienungsanleitung.docx (schlicht, serifenlos). Aufruf: python3 src/md2docx.py"""
import os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_BREAK

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
FONT = 'Helvetica Neue'

doc = Document()
for name in ('Normal', 'List Bullet', 'List Number'):
    st = doc.styles[name]; st.font.name = FONT; st.font.size = Pt(12)
for lvl, size in ((1, 20), (2, 15)):
    st = doc.styles['Heading %d' % lvl]
    st.font.name = FONT; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0x23, 0x46, 0x8f)

def runs(par, text):
    for i, part in enumerate(re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)):
        if not part: continue
        if part.startswith('**'): par.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Menlo'; r.font.size = Pt(10.5)
        else:
            par.add_run(re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', part))

lines = open(os.path.join(REPO, 'ANLEITUNG.md'), encoding='utf8').read().split('\n')
i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    s = ln.strip()
    if s.startswith('```'):
        i += 1; code = []
        while i < len(lines) and not lines[i].strip().startswith('```'):
            code.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); r = p.add_run('\n'.join(code)); r.font.name = 'Menlo'; r.font.size = Pt(9.5)
    elif s.startswith('# '): doc.add_heading(s[2:], 1)
    elif s.startswith('## '): doc.add_heading(s[3:], 2)
    elif re.match(r'^- ', s): runs(doc.add_paragraph(style='List Bullet'), s[2:])
    elif re.match(r'^\d+\. ', s): runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', s))
    elif s: runs(doc.add_paragraph(), s)
    i += 1
out = os.path.join(REPO, 'docs', 'Bedienungsanleitung.docx')
doc.save(out); print(out)
